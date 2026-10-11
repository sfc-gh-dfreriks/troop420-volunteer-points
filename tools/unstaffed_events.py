"""Troop 420: upcoming calendar activities with no adult in charge -> branded XLSX.

Reads TroopWebHost exports pulled with troop420-tools (`--preset volunteer-points`):
  53104 Export Events To Excel  (Type, Event, Start, Leaders, Sign-Up Deadline, Adults, Campout Nights, ...)
  45935 Export Event Participation, 54040 Adult Leaders By Position, 57107 Adult Involvement
        -> the list of known adults (anyone else in Leaders is treated as a Scout).
Leaders that are Scouts are counted, never named (youth names stay out of the file).

    python3 tools/unstaffed_events.py [--start 2026-10-11] [--months 6] [--exports ~/Documents/Troop420/exports]
"""
import argparse, csv, datetime as dt, glob, os, re
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

RED, GREEN, GOLD = "BB0000", "2D3D3C", "F4C72B"

# Points an adult lead would earn (mirrors config.js twhEventLeadTypes / rules.js)
LEAD_POINTS = {
    "Community Service": ("Service project lead", 25), "Conservation-related Service Project": ("Service project lead", 25),
    "Troop Service Day": ("Service project lead", 25), "Fundraiser": ("Fundraiser event lead", 25),
    "Tree Lot Fundraiser": ("Fundraiser shift", 10), "Court of Honor": ("Court of Honor coordination", 15),
    "Eagle Court of Honor": ("Court of Honor coordination", 15), "Merit Badge Clinic": ("Merit badge counselling", 15),
    "MB Day / Weekend": ("Merit badge counselling", 15), "Recruiting Event": ("Recruiting night volunteer", 10),
    "Hike": ("Day-only support", 10), "Outing / Field Trip": ("Day-only support", 10), "Cycling": ("Day-only support", 10),
    "Campout": ("Overnight adult (per trip)", 25), "Training / Campout": ("Overnight adult (per trip)", 25),
    "High Adventure Trek": ("High Adventure contingent lead", 60), "Setup / Teardown": ("Court of Honor setup", 10),
}
# Calendar entries that are not activities needing an adult in charge
NOT_ACTIVITY_TYPES = {"Holiday", "Advancement Deadline"}
NOT_ACTIVITY_NAME = re.compile(r"\bno (troop )?meeting\b|\bcancel+ed\b|\bdeadline\b|\bspring break\b|\bwinter break\b", re.I)
PLACEHOLDER = re.compile(r"^\s*(hold|tbd|tentative)\b", re.I)
# Adult-run by definition (the committee / SM run them) - shown, but lower priority
ADULT_RUN = {"Committee Meeting", "SM Conference", "Board of Review"}


def latest(d, mid):
    fs = sorted(glob.glob(os.path.join(d, f"report_{mid}_*.csv")))
    return list(csv.DictReader(open(fs[-1], encoding="utf-8-sig"))) if fs else []


def iso(s):
    m = re.match(r"(\d{1,2})/(\d{1,2})/(\d{2,4})", s or "")
    if not m: return None
    y = int(m.group(3)); y = y + 2000 if y < 100 else y
    return dt.date(y, int(m.group(1)), int(m.group(2)))


def time_of(s):
    m = re.search(r"(\d{1,2}:\d{2}\s*[AP]M)", s or "", re.I)
    return m.group(1).upper() if m else ""


def split_leaders(s):
    return [x.strip() for x in re.split(r"(?<=[a-z])(?=[A-Z][a-z]+\s)|\s*[;,\n]\s*", s or "") if x.strip()]


def adult_index(exports):
    """Known adults as (first-name, surname) keys, from every adult-bearing export."""
    names = set()
    for r in latest(exports, "45935"):
        if r.get("Adult?", "").upper().startswith("Y"): names.add(r["Participant"])
    for r in latest(exports, "54040"): names.add(r["Adult"])
    for r in latest(exports, "57107"): names.add(r["Name"])
    keys = set()
    for n in names:
        last, _, first = n.partition(",")
        if first.strip(): keys.add((first.strip().split()[0].lower(), last.strip().lower()))
    return keys


def is_adult(name, keys):
    parts = re.sub(r"[^A-Za-z\s'-]", "", name).lower().split()
    if len(parts) < 2: return False
    f, l = parts[0], parts[-1]
    return any(l == kl and (f == kf or (len(f) >= 3 and len(kf) >= 3 and (f.startswith(kf) or kf.startswith(f)))) for kf, kl in keys)


def build(exports, start, end, out):
    keys = adult_index(exports)
    rows, seen = [], set()
    for r in latest(exports, "53104"):
        d = iso(r["Start"])
        if not d or not (start <= d <= end): continue
        k = (r["Event"].strip(), r["Start"].strip())
        if k in seen: continue
        seen.add(k)
        leaders = split_leaders(r["Leaders"])
        adults = [x for x in leaders if is_adult(x, keys)]
        scouts = len(leaders) - len(adults)
        typ, name = r["Type"].strip(), r["Event"].strip()
        if typ in NOT_ACTIVITY_TYPES or NOT_ACTIVITY_NAME.search(name): status = "Not an activity"
        elif adults: status = "Adult lead"
        elif scouts: status = "Scout-led only - no adult"
        else: status = "No leader listed"
        role, pts = LEAD_POINTS.get(typ, ("", None))
        rows.append(dict(date=d, day=d.strftime("%a"), time=time_of(r["Start"]), type=typ, event=name,
                         location=r["Location"].strip(), status=status, adult_leads=", ".join(adults),
                         scout_leads=scouts, signup=iso(r["Sign-Up Deadline"]), adults_signed=int(float(r["Adults"] or 0)),
                         nights=int(float(r["Campout Nights"] or 0)), placeholder=bool(PLACEHOLDER.search(name)),
                         adult_run=typ in ADULT_RUN, role=role, pts=pts))
    rows.sort(key=lambda x: (x["date"], x["time"]))
    gaps = [x for x in rows if x["status"] in ("No leader listed", "Scout-led only - no adult")]

    def priority(x):
        if x["placeholder"]: return "Low - placeholder/hold"
        if x["adult_run"]: return "Low - committee/SM run"
        if x["nights"] or x["type"] in ("Campout", "Training / Campout", "High Adventure Trek"): return "High - overnight"
        if x["type"] in ("Tree Lot Fundraiser", "Fundraiser", "Community Service", "Conservation-related Service Project",
                         "Troop Service Day", "Court of Honor", "Eagle Court of Honor", "Recruiting Event"): return "High - needs a lead"
        return "Medium"
    for x in gaps: x["priority"] = priority(x)

    wb = Workbook()
    thin = Side(style="thin", color="D9D9D9")
    hdr_fill, hdr_font = PatternFill("solid", fgColor=GREEN), Font(bold=True, color="FFFFFF")
    pri_fill = {"High - overnight": "FDE2E2", "High - needs a lead": "FFF1CC", "Medium": "FFFFFF",
                "Low - committee/SM run": "EEF2F1", "Low - placeholder/hold": "EEF2F1"}

    def banner(ws, title, sub, ncols):
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
        ws["A1"] = title; ws["A1"].font = Font(bold=True, size=16, color="FFFFFF")
        ws["A1"].fill = PatternFill("solid", fgColor=RED); ws.row_dimensions[1].height = 28
        ws["A2"] = sub; ws["A2"].font = Font(italic=True, color=GREEN)
        ws["A2"].fill = PatternFill("solid", fgColor="FFF8DC")
        for c in range(1, ncols + 1): ws.cell(row=3, column=c).fill = PatternFill("solid", fgColor=GOLD)
        ws.row_dimensions[3].height = 4

    def table(ws, name, header, data, widths, fills=None, first=5):
        for i, h in enumerate(header, 1):
            c = ws.cell(row=first - 1, column=i, value=h); c.fill, c.font = hdr_fill, hdr_font
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for ri, rec in enumerate(data):
            for ci, v in enumerate(rec, 1):
                c = ws.cell(row=first + ri, column=ci, value=v)
                c.border = Border(bottom=thin); c.alignment = Alignment(vertical="top", wrap_text=ci in (5, 6, 14))
                if isinstance(v, dt.date): c.number_format = "ddd d mmm yyyy"
                if fills and fills[ri]: c.fill = PatternFill("solid", fgColor=fills[ri])
        for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
        if data:
            ref = f"A{first - 1}:{get_column_letter(len(header))}{first - 1 + len(data)}"
            t = Table(displayName=name, ref=ref)
            t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
            ws.add_table(t)
        ws.freeze_panes = ws.cell(row=first, column=3)
        ws.sheet_view.zoomScale = 110
        ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.print_title_rows = f"{first - 1}:{first - 1}"

    span = f"{start:%d %b %Y} - {end:%d %b %Y}"
    # 1. Needs an adult
    ws = wb.active; ws.title = "Needs an adult"
    H = ["Date", "Day", "Time", "Type", "Event", "Location", "Priority", "Why", "Scouts listed as leader",
         "Adults signed up", "Nights", "Sign-up deadline", "Lead credit (points list)", "Notes"]
    data, fills = [], []
    for x in gaps:
        data.append([x["date"], x["day"], x["time"], x["type"], x["event"], x["location"], x["priority"],
                     "Scout-led - no adult listed" if x["scout_leads"] else "No leader listed",
                     x["scout_leads"] or "", x["adults_signed"] or "", x["nights"] or "", x["signup"],
                     f'{x["role"]} - {x["pts"]} pts' if x["pts"] else "", ""])
        fills.append(pri_fill[x["priority"]])
    banner(ws, "Troop 420 Bothell - activities with no adult in charge",
           f"{span}  |  {len(gaps)} activities  |  from TroopWebHost calendar (Export Events To Excel), {dt.date.today():%d %b %Y}", len(H))
    table(ws, "NeedsAdult", H, data, [15, 6, 9, 20, 42, 26, 20, 24, 10, 10, 8, 15, 30, 30], fills)

    # 2. Summary
    s = wb.create_sheet("Summary")
    banner(s, "Summary - who will run it?", span, 6)
    acts = [x for x in rows if x["status"] != "Not an activity"]
    s["A5"], s["B5"] = "Calendar activities in window", len(acts)
    s["A6"], s["B6"] = "With an adult in charge", sum(x["status"] == "Adult lead" for x in acts)
    s["A7"], s["B7"] = "Scout-led only (no adult listed)", sum(x["status"].startswith("Scout") for x in acts)
    s["A8"], s["B8"] = "No leader listed", sum(x["status"] == "No leader listed" for x in acts)
    s["A9"], s["B9"] = "Not activities (holidays, no-meeting, deadlines)", len(rows) - len(acts)
    for r_ in range(5, 10): s.cell(row=r_, column=1).font = Font(bold=r_ in (7, 8)); s.cell(row=r_, column=2).font = Font(bold=True)
    pc = Counter(x["priority"] for x in gaps)
    r0 = 11
    for i, h in enumerate(["Priority", "Activities"], 1):
        c = s.cell(row=r0, column=i, value=h); c.fill, c.font = hdr_fill, hdr_font
    for i, p in enumerate(["High - overnight", "High - needs a lead", "Medium", "Low - committee/SM run", "Low - placeholder/hold"], 1):
        s.cell(row=r0 + i, column=1, value=p).fill = PatternFill("solid", fgColor=pri_fill[p]); s.cell(row=r0 + i, column=2, value=pc.get(p, 0))
    r1 = r0 + 8
    months = Counter(x["date"].strftime("%Y-%m") for x in gaps)
    types = Counter(x["type"] for x in gaps)
    for i, h in enumerate(["Month", "Needs an adult", "", "Type", "Needs an adult"], 1):
        if h: c = s.cell(row=r1, column=i, value=h); c.fill, c.font = hdr_fill, hdr_font
    for i, (m, n) in enumerate(sorted(months.items()), 1):
        s.cell(row=r1 + i, column=1, value=dt.datetime.strptime(m, "%Y-%m").strftime("%b %Y")); s.cell(row=r1 + i, column=2, value=n)
    for i, (t, n) in enumerate(types.most_common(), 1):
        s.cell(row=r1 + i, column=4, value=t); s.cell(row=r1 + i, column=5, value=n)
    for col, w in zip("ABCDE", (46, 16, 3, 34, 16)): s.column_dimensions[col].width = w
    s["A" + str(r1 + max(len(months), len(types)) + 2)] = (
        "Lead credit shows what an adult earns for running each activity on the Troop 420 Adult Volunteer Points list. "
        "Scout leaders are counted, never named. Source: TroopWebHost calendar; refresh with "
        "`python3 twh_export.py --login --preset volunteer-points` then `python3 tools/unstaffed_events.py`.")
    s["A" + str(r1 + max(len(months), len(types)) + 2)].alignment = Alignment(wrap_text=True)

    # 3. All activities
    a = wb.create_sheet("All calendar items")
    H2 = ["Date", "Day", "Time", "Type", "Event", "Location", "Status", "Adult lead(s)", "Scouts listed as leader",
          "Adults signed up", "Nights", "Sign-up deadline"]
    st_fill = {"Adult lead": "E6F4EA", "Scout-led only - no adult": "FFF1CC", "No leader listed": "FDE2E2", "Not an activity": "EEF2F1"}
    banner(a, "All calendar items", f"{span}  |  {len(rows)} entries", len(H2))
    table(a, "AllItems", H2, [[x["date"], x["day"], x["time"], x["type"], x["event"], x["location"], x["status"], x["adult_leads"],
                               x["scout_leads"] or "", x["adults_signed"] or "", x["nights"] or "", x["signup"]] for x in rows],
          [15, 6, 9, 20, 42, 26, 24, 30, 10, 10, 8, 15], [st_fill[x["status"]] for x in rows])
    wb.active = 0
    wb.save(out)
    return len(rows), len(gaps), pc


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--start", default=dt.date.today().isoformat())
    ap.add_argument("--months", type=int, default=6)
    ap.add_argument("--exports", default=os.path.expanduser("~/Documents/Troop420/exports"))
    ap.add_argument("--out", default=None)
    a = ap.parse_args()
    start = dt.date.fromisoformat(a.start)
    m = start.month - 1 + a.months
    end = dt.date(start.year + m // 12, m % 12 + 1, min(start.day, 28))
    out = a.out or os.path.expanduser(f"~/Documents/Troop420/reports/Troop420_Activities_Needing_An_Adult_{start:%Y%m%d}.xlsx")
    n, g, pc = build(a.exports, start, end, out)
    print(f"{out}\n{n} calendar items, {g} with no adult in charge: {dict(pc)}")
